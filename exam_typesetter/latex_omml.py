"""
LaTeX to Word OMML (Office Math ML) Converter
Supports inline formulas ($...$) and block formulas ($$...$$).
"""
import re
from docx.oxml import parse_xml
import latex2mathml.converter
import mathml2omml

NS_MATH = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

def latex_to_omml_xml(latex_str: str, is_display: bool = False) -> str:
    """
    Converts LaTeX math string to OMML XML string.
    """
    cleaned = latex_str.strip()
    # latex2mathml conversion
    mml = latex2mathml.converter.convert(cleaned)
    omml = mathml2omml.convert(mml)
    
    # Fix known mathml2omml typo for vectors/accents: </m:groupChr><m:e> -> </m:groupChrPr><m:e>
    if '</m:groupChr><m:e>' in omml:
        omml = omml.replace('</m:groupChr><m:e>', '</m:groupChrPr><m:e>')

    # Ensure math namespace is properly declared
    if '<m:oMath ' not in omml and '<m:oMath>' in omml:
        omml = omml.replace('<m:oMath>', f'<m:oMath {NS_MATH}>', 1)
        
    if is_display:
        return f'<m:oMathPara {NS_MATH}>{omml}</m:oMathPara>'
    return omml

def append_text_with_math(paragraph, text: str, font_name: str = "宋体", ascii_font: str = "Times New Roman", font_size_pt: float = 10.5):
    """
    Parses a string containing $inline$ and $$block$$ math expressions and appends them
    to a python-docx Paragraph with appropriate fonts and OMML math elements.
    """
    if not text:
        return

    # Pattern matches $$...$$ or $...$
    pattern = re.compile(r'(\$\$.*?\$\$|\$.*?\$)', re.DOTALL)
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue
            
        if part.startswith('$$') and part.endswith('$$'):
            # Block math
            latex = part[2:-2].strip()
            try:
                omml_xml = latex_to_omml_xml(latex, is_display=True)
                elem = parse_xml(omml_xml)
                paragraph._p.append(elem)
            except Exception:
                # Fallback to plain text on syntax error
                run = paragraph.add_run(part)
                _apply_font(run, font_name, ascii_font, font_size_pt)
                
        elif part.startswith('$') and part.endswith('$'):
            # Inline math
            latex = part[1:-1].strip()
            try:
                omml_xml = latex_to_omml_xml(latex, is_display=False)
                elem = parse_xml(omml_xml)
                paragraph._p.append(elem)
            except Exception:
                # Fallback to plain text on syntax error
                run = paragraph.add_run(part)
                _apply_font(run, font_name, ascii_font, font_size_pt)
                
        else:
            # Regular text
            run = paragraph.add_run(part)
            _apply_font(run, font_name, ascii_font, font_size_pt)

def _apply_font(run, east_asia_font: str, ascii_font: str, size_pt: float):
    run.font.name = ascii_font
    run.font.size = docx.shared.Pt(size_pt)
    # Set East Asian font in oxml element
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(docx.oxml.ns.qn('w:rFonts'))
    if rFonts is None:
        rFonts = docx.oxml.OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(docx.oxml.ns.qn('w:eastAsia'), east_asia_font)

import docx.shared
import docx.oxml
